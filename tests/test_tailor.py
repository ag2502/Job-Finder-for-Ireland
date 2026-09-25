"""Tailoring a CV to a job: the document engine, the guards, the score and the flow.

The model is stubbed throughout. What these tests hold is everything around it: that a
rewrite lands in the right place in every kind of file without disturbing the layout,
that nothing the rules forbid survives - an invented figure, a changed job title - and
that the searcher can review, revise, save, download or cancel.

The fixture CVs in tests/fixtures are synthetic: a fictional person, built with Chromium
and python-docx so their fonts and structure are like real exported CVs.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from jobfinder.tailor import ats, document, llm, proofread, rewrite

FIXTURES = Path(__file__).parent / "fixtures"


def _load(name: str) -> document.CvDocument:
    return document.load((FIXTURES / name).read_bytes(), name)


def _by_text(doc: document.CvDocument, start: str) -> document.Paragraph:
    return next(p for p in doc.paragraphs if p.text.startswith(start))


@pytest.fixture(autouse=True)
def _no_network(monkeypatch: pytest.MonkeyPatch):
    """No proofreading server in tests; the lint still runs."""
    monkeypatch.setattr(proofread.settings, "languagetool_url", "")


# ------------------------------------------------------------------ reading


@pytest.mark.parametrize("name", ["cv.pdf", "cv_two_column.pdf", "cv.docx", "cv.txt"])
def test_every_layout_reads_and_locks_the_facts(name):
    doc = _load(name)
    assert doc.paragraphs
    for p in doc.paragraphs:
        text = p.text
        if "@example.com" in text or "+353" in text:
            assert p.locked, f"contact details must be locked: {text}"
        if p.kind == "heading":
            assert p.locked
    assert any(not p.locked for p in doc.paragraphs), "something must be tailorable"


def test_dates_titles_and_education_are_locked():
    doc = _load("cv.pdf")
    assert _by_text(doc, "Mar 2021").locked
    assert _by_text(doc, "Fintech Co, Dublin").locked
    assert _by_text(doc, "MSc Data Analytics").locked
    assert not _by_text(doc, "Data scientist with five years").locked
    assert not _by_text(doc, "Tools:").locked, "a skills line is exactly what tailoring extends"


def test_bullets_drawn_as_shapes_are_separate_paragraphs():
    doc = _load("cv.pdf")
    bullets = [p for p in doc.paragraphs if p.kind == "bullet"]
    assert len(bullets) == 6
    assert bullets[0].text.endswith("40,000 customers.")


def test_a_two_column_cv_is_read_column_by_column():
    doc = _load("cv_two_column.pdf")
    order = [p.text for p in doc.paragraphs]
    assert order.index("SKILLS") < order.index("LANGUAGES") < order.index("SUMMARY")
    audited = _by_text(doc, "Audited financial statements")
    assert not audited.locked, "a main-column bullet is not part of the sidebar's languages"


def test_bold_phrases_are_marked_for_the_model():
    doc = _load("cv.docx")
    tools = _by_text(doc, "Tools:")
    assert tools.marked.startswith("**Tools:**")


# ------------------------------------------------------------------ writing


def test_a_pdf_rewrite_keeps_reading_order_and_reads_back_exactly():
    doc = _load("cv.pdf")
    profile = _by_text(doc, "Data scientist with five years")
    new = ("Data scientist with five years of experience building machine learning models "
           "and running experiments on product and payments data, from the first question "
           "to a production pipeline, explained clearly to non-technical stakeholders.")
    rendered = doc.render({profile.id: new})
    assert rendered.applied == {profile.id: new}
    text = document.text_of(rendered.data, "pdf")
    # In stored order, as a parser that does not sort by position would read it.
    assert text.index("PROFILE") < text.index("Data scientist") < text.index("EXPERIENCE")
    assert "­" not in text, "a hyphen must not read back as a soft hyphen"
    assert " ".join(new.split()) in " ".join(text.split())
    assert "turning messy product" not in text, "the old words are gone"


def test_a_pdf_rewrite_that_cannot_fit_is_left_as_it_was():
    doc = _load("cv.pdf")
    bullet = _by_text(doc, "Mentored two junior")
    rendered = doc.render({bullet.id: bullet.text + " " + bullet.text * 5})
    assert bullet.id in rendered.skipped
    assert bullet.text in document.text_of(rendered.data, "pdf")


def test_locked_paragraphs_are_never_rendered_even_if_asked():
    doc = _load("cv.pdf")
    name = doc.paragraphs[0]
    rendered = doc.render({name.id: "Somebody Else"})
    assert rendered.applied == {}


def test_a_word_rewrite_keeps_its_bold_lead_in():
    import docx

    doc = _load("cv.docx")
    tools = _by_text(doc, "Tools:")
    rendered = doc.render({tools.id: "**Tools:** Jupyter, scikit-learn, pandas, dbt"})
    out = docx.Document(io.BytesIO(rendered.data))
    para = next(p for p in out.paragraphs if p.text.startswith("Tools:"))
    assert para.text == "Tools: Jupyter, scikit-learn, pandas, dbt"
    assert para.runs[0].bold and para.runs[0].text.startswith("Tools:")
    assert not para.runs[-1].bold


def test_a_word_cv_can_lose_an_irrelevant_bullet():
    import docx

    doc = _load("cv.docx")
    bullet = _by_text(doc, "Wrote SQL and dbt")
    rendered = doc.render({bullet.id: ""})
    texts = [p.text for p in docx.Document(io.BytesIO(rendered.data)).paragraphs]
    assert not any(t.startswith("Wrote SQL") for t in texts)


def test_a_pdf_never_loses_a_paragraph():
    doc = _load("cv.pdf")
    bullet = _by_text(doc, "Mentored two junior")
    rendered = doc.render({bullet.id: ""})
    assert rendered.applied == {}
    assert bullet.text in document.text_of(rendered.data, "pdf")


# ------------------------------------------------------------------ the guards


def _answer(edits, **extra):
    return {"keywords": ["machine learning", "A/B testing", "Python"], "edits": edits,
            "fit_summary": "It fits.", "strengths": ["Python"], "gaps": ["Spark"], **extra}


JOB = rewrite.Job("Senior Data Scientist, Payments", "Stripe",
                  "We want machine learning, A/B testing, Python, SQL and Spark. " * 8)


def test_an_invented_figure_or_a_locked_line_is_refused(monkeypatch):
    doc = _load("cv.pdf")
    bullet = _by_text(doc, "Built churn prediction")
    profile = _by_text(doc, "Data scientist with five years")
    monkeypatch.setattr(llm, "ask", lambda *a, **k: (_answer([
        {"id": bullet.id, "text": bullet.text.replace("12%", "15%"), "reason": "x"},
        {"id": doc.paragraphs[0].id, "text": "Someone Else", "reason": "x"},
        {"id": profile.id, "text": profile.text.replace("turning", "turning the the"), "reason": "y"},
    ]), "stub"))
    result = rewrite.tailor(doc, JOB)
    assert bullet.id not in result.edits
    assert doc.paragraphs[0].id not in result.edits
    assert "the the" not in result.edits[profile.id], "doubled words are linted out"
    whys = " ".join(d["why"] for d in result.report["dropped"])
    assert "15" in whys and "kept exactly as it is" in whys


def test_a_figure_the_candidate_gives_is_allowed_in_a_revision(monkeypatch):
    doc = _load("cv.pdf")
    bullet = _by_text(doc, "Wrote SQL and dbt")
    new = "Wrote SQL and dbt pipelines on Snowflake serving 25 revenue dashboards in Looker."
    monkeypatch.setattr(llm, "ask", lambda *a, **k: (_answer(
        [{"id": bullet.id, "text": new, "reason": "your figure"}], reply="Added the 25."), "stub"))
    result = rewrite.revise(doc, JOB, {}, {}, "It fed 25 dashboards.", ["Python"])
    assert result.edits[bullet.id] == new
    assert result.report["reply"] == "Added the 25."


def test_removal_is_capped_at_a_third_of_the_bullets(monkeypatch):
    doc = _load("cv.docx")
    bullets = [p for p in doc.paragraphs if p.kind == "bullet"]
    monkeypatch.setattr(llm, "ask", lambda *a, **k: (_answer(
        [{"id": b.id, "text": "", "reason": "off topic"} for b in bullets]), "stub"))
    result = rewrite.tailor(doc, JOB)
    removed = [pid for pid, text in result.edits.items() if not text]
    assert len(removed) == len(bullets) // 3


def test_no_models_means_a_plain_message(monkeypatch):
    def down(*a, **k):
        raise llm.ModelsUnavailable("all busy")

    monkeypatch.setattr(llm, "ask", down)
    with pytest.raises(rewrite.TailorError, match="busy"):
        rewrite.tailor(_load("cv.pdf"), JOB)


# ------------------------------------------------------------------ the score


def test_the_ats_score_is_deterministic_and_explained():
    text = document.text_of((FIXTURES / "cv.pdf").read_bytes(), "pdf")
    kwargs = dict(text=text, headings=["PROFILE", "EXPERIENCE", "EDUCATION", "SKILLS"],
                  bullets=["Cut churn by 12%.", "Led reviews."], keywords=["Python", "Spark"],
                  job_title="Senior Data Scientist, Payments", kind="pdf")
    first, second = ats.score(**kwargs), ats.score(**kwargs)
    assert first == second
    assert first["matched"] == ["Python"] and first["missing"] == ["Spark"]
    title = next(c for c in first["checks"] if c["label"] == "Job title")
    assert title["points"] == 10, "'data scientist' appears, which is the title without its level"
    assert all(c["detail"] for c in first["checks"])


def test_keywords_match_whole_words_only():
    assert ats._has("sql", "wrote sql pipelines")
    assert not ats._has("r", "worked with dbt and airflow")
    assert ats._has("dashboard", "built dashboards")


def test_the_advert_is_cut_to_its_requirements_when_long():
    text = "About us. " * 900 + "Requirements: Python and SQL."
    assert "Requirements: Python" in rewrite.job_excerpt(text)
